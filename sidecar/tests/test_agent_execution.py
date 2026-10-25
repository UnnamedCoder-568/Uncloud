"""Real local workflows, malformed plans, and failure/consent boundaries."""
from __future__ import annotations

import asyncio
import json
import os
import sys
from types import SimpleNamespace

import pytest

from uncloud_engine.agent import approval, execution, orchestrator, tools
from uncloud_engine.agent.contracts import tool_contract
from uncloud_engine.core import Denied, Gate


@pytest.fixture
def local(tmp_path, monkeypatch):
    monkeypatch.setattr(tools, 'WORKSPACE_DIR', tmp_path)
    monkeypatch.setattr(tools.settings, '_data', {'agent_device_access': True})
    monkeypatch.setattr(orchestrator.Orchestrator, '_persist_plan', lambda *args: None)
    async def completed_review(self, goal, **kwargs):
        assert kwargs.get('allow_empty'), 'Unexpected unmocked continuation'
        return orchestrator.ExecutionGraph(goal=goal)
    monkeypatch.setattr(orchestrator.Orchestrator, 'plan', completed_review)
    tools.install_gate(Gate())
    yield tmp_path
    tools.install_gate(None)


async def approve(request):
    return 'yes'


def run(tool, args):
    async def exercise():
        with approval.asking(approve):
            return await tools.run_tool(tool, args)
    return asyncio.run(exercise())


def test_shell_and_files_share_the_engine_python_and_workspace(local):
    assert json.loads(run('runtime_info', {}))['python'] == sys.executable
    run('shell', {'command': 'python -c "from pathlib import Path; '
                                'Path(\'unicode.txt\').write_text(\'hello\')"'})
    assert run('fs_read', {'path': 'unicode.txt'}) == 'hello'
    out = run('fs_write', {'path': 'unicode.txt', 'content': 'café'})
    assert '5 bytes' in out
    assert str(local) in out
    run('fs_edit', {'path': 'unicode.txt', 'old': 'café', 'new': 'edited'})
    assert 'unicode.txt' in run('fs_glob', {'pattern': '*.txt'})
    assert 'edited' in run('fs_grep', {'pattern': 'edited'})
    assert json.loads(run('file_info', {'path': 'unicode.txt'}))['bytes'] == 6


def test_file_tools_use_explicit_working_folder_and_preview_resolved_path(local):
    folder = local / 'nested folder'
    folder.mkdir()
    requested = []

    async def inspect(request):
        requested.append(request)
        return 'yes'

    async def exercise():
        with approval.asking(inspect):
            await tools.run_tool('fs_write', {'cwd': str(folder), 'path': 'data.txt',
                                              'content': 'café'})
    asyncio.run(exercise())
    assert requested[0].preview['path'] == str(folder / 'data.txt')
    assert requested[0].preview['bytes'] == 5
    assert run('fs_read', {'cwd': str(folder), 'path': 'data.txt'}) == 'café'
    run('fs_edit', {'cwd': str(folder), 'path': 'data.txt', 'old': 'café', 'new': 'edited'})
    assert 'data.txt' in run('fs_list', {'cwd': str(folder)})
    assert 'data.txt' in run('fs_glob', {'cwd': str(folder), 'pattern': '*.txt'})
    assert 'edited' in run('fs_grep', {'cwd': str(folder), 'pattern': 'edited'})
    assert json.loads(run('file_info', {'cwd': str(folder), 'path': 'data.txt'}))['bytes'] == 6
    assert not (local / 'data.txt').exists()


def test_file_cwd_cannot_escape_workspace_before_approval(local, monkeypatch):
    monkeypatch.setattr(tools.settings, '_data', {'agent_device_access': False})

    async def forbidden(*args, **kwargs):
        pytest.fail('Outside-workspace call reached approval')
    monkeypatch.setattr(approval, 'decide', forbidden)
    with pytest.raises(PermissionError, match='outside the agent workspace'):
        run('fs_write', {'cwd': str(local.parent), 'path': 'outside.txt', 'content': 'no'})
    with pytest.raises(PermissionError, match='outside the agent workspace'):
        run('fs_write', {'cwd': str(local), 'path': '../outside.txt', 'content': 'no'})


def test_python_can_return_stdout_without_declaring_a_file(local):
    result = json.loads(run('python_exec', {'code': 'print(7 + 3)'}))
    assert result == {'stdout': '10\n', 'outputs': []}


@pytest.mark.parametrize('tool', ['shell', 'python_exec'])
def test_code_approval_shows_the_actual_resolved_working_folder(local, tool):
    folder = local / 'working folder'
    folder.mkdir()
    requests = []

    async def ask(request):
        requests.append(request)
        return 'yes'

    async def exercise():
        args = {'cwd': 'working folder'}
        if tool == 'shell':
            args['command'] = 'python -c "import os; print(os.getcwd())"'
        else:
            args['code'] = 'import os; print(os.getcwd())'
        with approval.asking(ask):
            return await tools.run_tool(tool, args)
    result = asyncio.run(exercise())
    assert requests[0].preview['cwd'] == str(folder)
    output = json.loads(result)['stdout'] if tool == 'python_exec' else result
    assert output.strip() == str(folder)


def test_missing_declared_output_preserves_successful_stdout_for_recovery(local):
    with pytest.raises(ValueError, match=r'(?s)Output verification failed:.*Process stdout:\n10'):
        run('python_exec', {'code': 'print(7 + 3)', 'output_paths': ['never-written.txt']})


@pytest.mark.parametrize('suffix', ['pdf', 'docx', 'xlsx', 'pptx', 'png'])
def test_text_is_never_reported_as_a_valid_binary_document(local, suffix):
    with pytest.raises(ValueError, match='UTF-8 text'):
        run('fs_write', {'path': f'fake.{suffix}', 'content': 'not a document'})
    assert not (local / f'fake.{suffix}').exists()


def test_python_creates_readable_pdf_word_and_spreadsheet(local):
    from docx import Document
    from openpyxl import load_workbook
    from pypdf import PdfReader

    result = json.loads(run('python_exec', {'code': '''
from reportlab.pdfgen.canvas import Canvas
from docx import Document
from openpyxl import Workbook
c = Canvas('report.pdf'); c.drawString(60, 750, 'Chisel verified output'); c.save()
d = Document(); d.add_paragraph('Chisel verified output'); d.save('report.docx')
w = Workbook(); w.active.append(['quantity', 7]); w.save('report.xlsx')
''', 'output_paths': ['report.pdf', 'report.docx', 'report.xlsx']}))
    assert len(result['outputs']) == 3
    assert 'Chisel verified output' in PdfReader(local / 'report.pdf').pages[0].extract_text()
    assert Document(local / 'report.docx').paragraphs[0].text == 'Chisel verified output'
    assert load_workbook(local / 'report.xlsx').active['B1'].value == 7
    assert 'Chisel verified output' in run('fs_read', {'path': 'report.pdf'})
    assert 'Chisel verified output' in run('fs_read', {'path': 'report.docx'})
    assert 'quantity\t7' in run('fs_read', {'path': 'report.xlsx'})


@pytest.mark.parametrize('suffix', ['pdf', 'docx', 'xlsx'])
def test_structured_documents_save_and_read_back_real_content(local, suffix):
    result = json.loads(run('document_write', {'path': f'inventory.{suffix}',
        'title': 'Inventory', 'paragraphs': ['Checked inventory <literal> & café'],
        'rows': [['Item', 'Quantity'], ['Apples', 7], ['Oranges', 3], ['Total', 10]]}))
    assert result['path'] == str(local / f'inventory.{suffix}')
    assert result['bytes'] > 0 and result['format'] == suffix
    for text in ['Item', 'Quantity', 'Apples', 'Oranges', 'Total', '10', '<literal>', 'café']:
        assert text in result['text']


def test_structured_document_strings_are_data_not_spreadsheet_formulas(local):
    from openpyxl import load_workbook

    run('document_write', {'path': 'strings.xlsx', 'rows': [['=1+1', None, True, 3.5]]})
    workbook = load_workbook(local / 'strings.xlsx')
    try:
        assert workbook.active['A1'].value == '=1+1'
        assert workbook.active['A1'].data_type == 's'
        assert workbook.active['C1'].value is True
        assert workbook.active['D1'].value == 3.5
    finally:
        workbook.close()


def test_document_approval_shows_canonical_path_and_denial_writes_nothing(local):
    requested = []

    async def deny(request):
        requested.append(request)
        return 'no'

    async def exercise():
        with approval.asking(deny):
            await tools.run_tool('document_write', {'path': 'report.pdf',
                                                    'paragraphs': ['private']})
    with pytest.raises(Denied):
        asyncio.run(exercise())
    assert requested[0].preview['path'] == str(local / 'report.pdf')
    assert requested[0].preview['paragraphs'] == ['private']
    assert not (local / 'report.pdf').exists()


def test_document_tool_does_not_escape_confined_workspace(local, monkeypatch):
    monkeypatch.setattr(tools.settings, '_data', {'agent_device_access': False})
    with pytest.raises(PermissionError, match='outside the agent workspace'):
        run('document_write', {'path': str(local.parent / 'outside.pdf'), 'paragraphs': ['text']})
    result = json.loads(run('document_write', {'path': 'inside.pdf', 'paragraphs': ['Inside']}))
    assert result['text'].strip() == 'Inside'


def test_legacy_plan_can_use_observed_rows_for_a_document(local):
    (local / 'rows.json').write_text('[["Item","Quantity"],["Apples",7]]')
    graph = orchestrator._graph_from_plan({'tasks': {
        'read': {'tool_id': 'fs_read', 'args': {'path': 'rows.json'}},
        'write': {'tool_id': 'document_write', 'args': {'path': 'result.xlsx', 'rows': '{{read}}'},
                  'dependencies': ['read']},
    }}, 'Create spreadsheet from saved rows', tools.TOOL_SPECS)

    async def exercise():
        async def update(current):
            pass
        with approval.asking(approve):
            return await orchestrator.Orchestrator()._execute(graph, update)
    result = asyncio.run(exercise())
    assert not result.failed_tasks()
    assert 'Apples\t7' in json.loads(result.tasks['write'].output)['text']


def test_failed_document_render_preserves_existing_file_and_removes_temporary_output(
    local, monkeypatch,
):
    from reportlab.platypus import SimpleDocTemplate

    target = local / 'report.pdf'
    target.write_bytes(b'previous file')
    monkeypatch.setattr(SimpleDocTemplate, 'build', lambda *a, **k: (_ for _ in ()).throw(
        RuntimeError('Render failed')))
    with pytest.raises(RuntimeError, match='Render failed'):
        run('document_write', {'path': 'report.pdf', 'paragraphs': ['new content']})
    assert target.read_bytes() == b'previous file'
    assert not list(local.glob('.uncloud-document-*'))


@pytest.mark.parametrize('args', [
    {'path': 'wrong.txt', 'paragraphs': ['text']}, {'path': 'empty.pdf'},
    {'path': 'bad.pdf', 'rows': [['cell', {}]]},
])
def test_invalid_document_content_is_rejected_before_approval(local, monkeypatch, args):
    async def forbidden(*a, **k):
        pytest.fail('Invalid document reached approval')
    monkeypatch.setattr(approval, 'decide', forbidden)
    with pytest.raises(ValueError):
        run('document_write', args)


@pytest.mark.parametrize('code,outputs,error', [
    ('raise ValueError("deliberate failure")', [], 'Exit'),
    ('print("done")', ['missing.txt'], 'No such file'),
    ('from pathlib import Path; Path("fake.pdf").write_text("invalid")', ['fake.pdf'], ''),
    ('print("done")', ['existing.txt'], 'not updated'),
])
def test_code_errors_and_invalid_outputs_fail_the_graph(local, code, outputs, error):
    (local / 'existing.txt').write_text('old')
    graph = orchestrator._graph_from_plan({'tasks': {
        'create': {'tool_id': 'python_exec', 'args': {'code': code, 'output_paths': outputs}},
        'next': {'tool_id': 'fs_write', 'args': {'path': 'should-not-exist', 'content': 'bad'},
                 'dependencies': ['create']},
    }}, 'Create and use output', tools.TOOL_SPECS)

    async def exercise():
        async def update(graph):
            pass
        with approval.asking(approve):
            await orchestrator.Orchestrator()._execute(graph, update)
    asyncio.run(exercise())
    assert graph.tasks['create'].status == graph.tasks['next'].status == 'failed'
    assert error in graph.tasks['create'].error
    assert not (local / 'should-not-exist').exists()


@pytest.mark.parametrize('args', [{}, {'code': 123, 'output_paths': []},
                                 {'code': 'pass', 'output_paths': 'file.pdf'},
                                 {'code': 'pass', 'output_paths': [], 'unknown': True},
                                 {'code': 'if :', 'output_paths': []}])
def test_invalid_arguments_are_rejected_before_permission_or_execution(local, monkeypatch, args):
    async def forbidden(*args, **kwargs):
        pytest.fail('Invalid call reached approval or execution')
    monkeypatch.setattr(approval, 'decide', forbidden)
    monkeypatch.setattr(execution, 'execute', forbidden)
    with pytest.raises(ValueError):
        run('python_exec', args)


def test_python_remains_blocked_in_workspace_mode(local, monkeypatch):
    monkeypatch.setattr(tools.settings, '_data', {'agent_device_access': False})
    with pytest.raises(PermissionError, match='full device'):
        run('python_exec', {'code': 'print(1)', 'output_paths': []})


def test_python_approval_shows_code_and_can_be_refused(local):
    async def refuse(request):
        assert request.action == 'python_exec'
        assert 'should-not-exist' in request.preview['code']
        return 'no'

    async def exercise():
        with approval.asking(refuse):
            await tools.run_tool('python_exec', {
                'code': 'open("should-not-exist", "w").write("bad")', 'output_paths': []})
    with pytest.raises(Denied):
        asyncio.run(exercise())
    assert not (local / 'should-not-exist').exists()


@pytest.mark.parametrize('answer,reason', [('no', 'the user declined'),
                                         ('expired', 'expired without an answer')])
def test_refused_or_expired_approval_stops_all_actions_and_recovery(local, monkeypatch,
                                                                  answer, reason):
    graph = orchestrator._graph_from_plan({'tasks': {
        'first': {'tool_id': 'python_exec', 'args': {'code': 'print(1)'}},
        'independent': {'tool_id': 'fs_write', 'args': {'path': 'wrong.txt', 'content': 'no'}},
    }}, 'Permission boundary', tools.TOOL_SPECS)
    planner = orchestrator.Orchestrator()
    requests = []

    async def refuse(request):
        requests.append(request.action)
        return answer

    async def forbidden(*args, **kwargs):
        pytest.fail('A refused action must not trigger recovery')

    async def update(current):
        pass

    monkeypatch.setattr(planner, '_replan', forbidden)

    async def exercise():
        with approval.asking(refuse):
            return await planner.run(graph, update, effort='maximum')
    result = asyncio.run(exercise())
    assert requests == ['python_exec']
    assert reason in result.halt_reason
    assert result.to_dict()['halt_reason'] == result.halt_reason
    assert all(task.status == 'failed' for task in result.tasks.values())
    assert not (local / 'wrong.txt').exists()


def test_refusal_during_recovery_also_stops_later_recovery_passes(local, monkeypatch):
    graph = orchestrator._graph_from_plan({'tasks': {
        'read': {'tool_id': 'fs_read', 'args': {'path': 'missing.txt'}},
    }}, 'Recover missing input', tools.TOOL_SPECS)
    planner = orchestrator.Orchestrator()
    retries = []

    async def replan(current, failed, attempt):
        retries.append(attempt)
        return orchestrator._graph_from_plan({'tasks': {
            'retry': {'tool_id': 'python_exec', 'args': {'code': 'print(1)'}},
            'later': {'tool_id': 'fs_write', 'args': {'path': 'wrong.txt', 'content': 'no'}},
        }}, current.goal, tools.TOOL_SPECS)

    async def ask(request):
        return 'no' if request.action == 'python_exec' else 'yes'

    async def update(current):
        pass

    monkeypatch.setattr(planner, '_replan', replan)

    async def exercise():
        with approval.asking(ask):
            return await planner.run(graph, update, effort='maximum')
    result = asyncio.run(exercise())
    assert retries == [1]
    assert result.halt_reason == 'the user declined'
    assert result.tasks['later'].status == 'failed'
    assert not (local / 'wrong.txt').exists()


def test_code_stdout_reaches_the_model_before_a_later_action(local, monkeypatch):
    graph = orchestrator._graph_from_plan({'tasks': {
        'calculate': {'tool_id': 'python_exec', 'args': {'code': 'print(7 + 3)'}},
        'write': {'tool_id': 'fs_write', 'args': {'path': 'result.txt', 'content': 'stale'},
                  'dependencies': ['calculate']},
    }}, 'Save the calculated total', tools.TOOL_SPECS)
    planner = orchestrator.Orchestrator()
    observations = []

    async def plan(goal, **kwargs):
        if kwargs.get('allow_empty'):
            return orchestrator.ExecutionGraph(goal=goal)
        context = json.loads(goal)
        observed = json.loads(context['observations'][0]['output'])
        observations.append(observed['stdout'])
        return orchestrator._graph_from_plan({'tasks': {
            'write': {'tool_id': 'fs_write', 'args': {'path': 'result.txt',
                                                     'content': observed['stdout']}},
        }}, goal, tools.TOOL_SPECS, kwargs.get('completed'))

    async def update(current):
        pass

    monkeypatch.setattr(planner, 'plan', plan)

    async def exercise():
        with approval.asking(approve):
            return await planner.run(graph, update)
    result = asyncio.run(exercise())
    assert observations == ['10\n']
    assert (local / 'result.txt').read_text() == '10\n'
    assert not result.failed_tasks()


def test_execution_failure_stops_draft_actions_before_recovery(local):
    graph = orchestrator._graph_from_plan({'tasks': {
        'fail': {'tool_id': 'python_exec', 'args': {'code': 'raise RuntimeError("failed")'}},
        'draft': {'tool_id': 'fs_write', 'args': {'path': 'wrong.txt', 'content': 'no'}},
    }}, 'Stop stale actions', tools.TOOL_SPECS)

    async def update(current):
        pass

    async def exercise():
        with approval.asking(approve):
            return await orchestrator.Orchestrator().run(graph, update, effort='fast')
    result = asyncio.run(exercise())
    assert result.tasks['draft'].status == 'failed'
    assert 'Awaiting a revised plan' in result.tasks['draft'].error
    assert not (local / 'wrong.txt').exists()


def test_unknown_or_unordered_references_cannot_enter_a_plan():
    for reference in ('missing', 'read'):
        with pytest.raises(ValueError, match='dependency'):
            orchestrator._graph_from_plan({'tasks': {
                'read': {'tool_id': 'fs_read', 'args': {'path': 'input.txt'}},
                'write': {'tool_id': 'fs_write',
                          'args': {'path': 'out.txt', 'content': '{{' + reference + '}}'}},
            }}, 'goal', tools.TOOL_SPECS)


def test_reference_substitution_preserves_code_braces_and_uses_completed_results():
    graph = orchestrator._graph_from_plan({'tasks': {
        'read': {'tool_id': 'fs_read', 'args': {'path': 'input.txt'}},
        'write': {'tool_id': 'fs_write', 'dependencies': ['read'],
                  'args': {'path': 'out.py', 'content': 'f"{variable}" + {{read}}'}},
    }}, 'goal', tools.TOOL_SPECS)
    with pytest.raises(ValueError, match='not completed'):
        orchestrator._resolve_refs(graph.tasks['write'].args, graph)
    graph.tasks['read'].status = 'completed'
    graph.tasks['read'].output = 'actual result'
    assert orchestrator._resolve_refs(graph.tasks['write'].args, graph)['content'] == (
        'f"{variable}" + actual result')


def test_numeric_step_ids_from_native_model_plans_are_accepted():
    graph = orchestrator._graph_from_plan({'tasks': {
        '1': {'tool_id': 'fs_read', 'args': {'path': 'input.txt'}},
        '2': {'tool_id': 'fs_write', 'args': {'path': 'output.txt', 'content': '{{1}}'},
              'dependencies': ['1']},
    }}, 'Copy a file', tools.TOOL_SPECS)
    assert graph.ready_tasks()[0].id == '1'


def test_native_single_tool_object_is_validated_as_a_one_step_plan():
    value = {'id': 'create', 'tool_id': 'fs_write',
             'args': {'path': 'output.py', 'content': 'print(7)'}}
    graph = orchestrator._graph_from_plan(orchestrator._extract_json(json.dumps(value)),
                                          'Create a script', tools.TOOL_SPECS)
    assert graph.tasks['create'].tool_id == 'fs_write'


def test_structured_tool_output_can_feed_typed_arguments():
    graph = orchestrator._graph_from_plan({'tasks': {
        'read': {'tool_id': 'fs_read', 'args': {'path': 'arguments.json'}},
        'use': {'tool_id': 'capability', 'args': {
            'capability': 'document.create', 'arguments': '{{read}}'}, 'dependencies': ['read']},
    }}, 'Use saved arguments', tools.TOOL_SPECS)
    graph.tasks['read'].status = 'completed'
    graph.tasks['read'].output = '{"title":"Report"}'
    resolved = orchestrator._resolve_refs(graph.tasks['use'].args, graph, 'capability')
    assert resolved['arguments'] == {'title': 'Report'}


def test_new_reply_labels_cannot_overwrite_completed_task_results():
    first = orchestrator._graph_from_plan({'tasks': {
        't1': {'tool_id': 'fs_read', 'args': {'path': 'first.txt'}},
    }}, 'First read', tools.TOOL_SPECS)
    first.tasks['t1'].status = 'completed'
    first.tasks['t1'].output = 'preserved earlier result'
    next_plan = orchestrator._graph_from_plan({'tasks': {
        't1': {'tool_id': 'fs_read', 'args': {'path': 'second.txt'}},
        't2': {'tool_id': 'fs_write', 'args': {'path': 'result.txt', 'content': '{{t1}}'},
               'dependencies': ['t1']},
    }}, 'Use a new input', tools.TOOL_SPECS, first.tasks)
    assert set(next_plan.tasks) == {'next_t1', 't2'}
    assert next_plan.tasks['t2'].dependencies == ['next_t1']
    assert next_plan.tasks['t2'].args['content'] == '{{next_t1}}'
    assert first.tasks['t1'].output == 'preserved earlier result'
    next_plan.tasks['next_t1'].status = 'completed'
    next_plan.tasks['next_t1'].output = 'new observed result'
    assert orchestrator._resolve_refs(next_plan.tasks['t2'].args, next_plan,
                                       'fs_write')['content'] == 'new observed result'


def test_reply_label_renaming_preserves_python_variables_with_matching_names():
    renamed = orchestrator._rename_refs({
        'code': 't1 = 7\nprint(f"{t1}")\nprint({"t1": t1})',
        'result': '{{t1}}',
        'legacy_result': '{t1}',
    }, {'t1': 'next_t1'})
    assert renamed['code'] == 't1 = 7\nprint(f"{t1}")\nprint({"t1": t1})'
    assert renamed['result'] == '{{next_t1}}'
    assert renamed['legacy_result'] == '{next_t1}'


def test_completed_action_labels_retain_results_instead_of_repeating_side_effects():
    completed = {
        'TASK_1': orchestrator.Task('TASK_1', 'Write', 'fs_write',
            {'path': 'notes.txt', 'content': 'alpha'}, status='completed', output='written'),
        'TASK_2': orchestrator.Task('TASK_2', 'Edit', 'fs_edit',
            {'path': 'notes.txt', 'old': 'alpha', 'new': 'beta'},
            dependencies=['TASK_1'], status='completed', output='edited'),
    }
    plan = {'tasks': {ident: {
        'tool_id': task.tool_id, 'args': task.args, 'dependencies': task.dependencies,
    } for ident, task in completed.items()}}
    with pytest.raises(ValueError, match='only repeats completed'):
        orchestrator._graph_from_plan(plan, 'Continue', tools.TOOL_SPECS, completed)
    plan['tasks']['TASK_3'] = {'tool_id': 'fs_read', 'args': {'path': 'notes.txt'},
                               'dependencies': ['TASK_2']}
    result = orchestrator._graph_from_plan(plan, 'Continue', tools.TOOL_SPECS, completed)
    assert set(result.tasks) == {'TASK_3'}
    assert result.tasks['TASK_3'].dependencies == ['TASK_2']
    assert completed['TASK_2'].output == 'edited'


def test_changed_local_input_prevents_reusing_a_completed_dependent_action():
    completed = {
        'read': orchestrator.Task('read', 'Read', 'fs_read', {'path': 'first.txt'},
                                  status='completed', output='first'),
        'write': orchestrator.Task('write', 'Write', 'fs_write',
            {'path': 'result.txt', 'content': '{{read}}'}, dependencies=['read'],
            status='completed', output='written'),
    }
    result = orchestrator._graph_from_plan({'tasks': {
        'read': {'tool_id': 'fs_read', 'args': {'path': 'second.txt'}},
        'write': {'tool_id': 'fs_write', 'args': completed['write'].args,
                  'dependencies': ['read']},
    }}, 'Use the second input', tools.TOOL_SPECS, completed)
    assert set(result.tasks) == {'next_read', 'next_write'}
    assert result.tasks['next_write'].args['content'] == '{{next_read}}'
    assert completed['read'].output == 'first'


def test_launcher_failures_are_tool_failures_not_successful_text(monkeypatch):
    import subprocess

    monkeypatch.setattr(sys, 'platform', 'darwin')
    monkeypatch.setattr(subprocess, 'run', lambda *a, **k: SimpleNamespace(
        returncode=1, stderr='Application not found'))
    with pytest.raises(RuntimeError, match='Application not found'):
        tools._app_open('', 'Missing App')


def test_windows_launcher_does_not_interpret_url_metacharacters_as_commands(monkeypatch):
    seen = []
    monkeypatch.setattr(sys, 'platform', 'win32')
    monkeypatch.setattr(os, 'startfile', seen.append, raising=False)
    target = 'https://example.test/?one=1&two=2'
    assert 'Opened' in tools._app_open(target)
    assert seen == [target]


def test_successful_recovery_stops_retries_and_preserves_the_original_error(local, monkeypatch):
    original = orchestrator._graph_from_plan({'tasks': {
        '1': {'tool_id': 'fs_read', 'args': {'path': 'missing.txt'}},
    }}, 'Create a file', tools.TOOL_SPECS)
    recovered = orchestrator._graph_from_plan({'tasks': {
        'retry': {'tool_id': 'fs_write', 'args': {'path': 'fixed.txt', 'content': 'ok'}},
    }}, 'Create a file', tools.TOOL_SPECS)
    planner = orchestrator.Orchestrator()
    retries = []

    async def replan(graph, failed, attempt):
        retries.append(attempt)
        return recovered

    monkeypatch.setattr(planner, '_replan', replan)

    async def exercise():
        async def update(graph):
            pass
        with approval.asking(approve):
            return await planner.run(original, update, effort='maximum')

    graph = asyncio.run(exercise())
    assert retries == [1]
    assert not graph.failed_tasks()
    assert graph.tasks['1'].error
    assert graph.tasks['1'].recovered_by == ['retry']
    assert (local / 'fixed.txt').read_text() == 'ok'


def test_recovery_planning_error_is_visible_in_final_progress(local, monkeypatch):
    graph = orchestrator._graph_from_plan({'tasks': {
        'read': {'tool_id': 'fs_read', 'args': {'path': 'missing.txt'}},
    }}, 'Read missing input', tools.TOOL_SPECS)
    planner = orchestrator.Orchestrator()
    updates = []

    async def plan(*args, **kwargs):
        raise RuntimeError('Invalid recovery arguments')

    async def update(current):
        updates.append(current.tasks['read'].error)

    monkeypatch.setattr(planner, 'plan', plan)

    async def exercise():
        with approval.asking(approve):
            return await planner.run(graph, update)
    result = asyncio.run(exercise())
    assert result.failed_tasks()
    assert 'No such file' in updates[-1]
    assert 'Recovery planning failed: Invalid recovery arguments' in updates[-1]


def test_observations_update_future_actions_before_any_draft_write(local, monkeypatch):
    (local / 'input.json').write_text('{"quantity":7}')
    original = orchestrator._graph_from_plan({'tasks': {
        'read': {'description': 'Read input', 'tool_id': 'fs_read',
                 'args': {'path': 'input.json'}},
        'write': {'description': 'Write the observed quantity', 'tool_id': 'fs_write',
                  'args': {'path': 'result.txt', 'content': ''}},
    }}, 'Write the quantity from input.json', tools.TOOL_SPECS)
    planner = orchestrator.Orchestrator()
    calls = []

    async def plan(goal, **kwargs):
        if kwargs.get('allow_empty'):
            return orchestrator.ExecutionGraph(goal=goal)
        request = json.loads(goal)
        calls.append(request)
        assert not (local / 'result.txt').exists()
        assert request['observations'][0]['output'] == '{"quantity":7}'
        assert request['remaining'][0]['args'] == {'path': 'result.txt', 'content': ''}
        quantity = json.loads(request['observations'][0]['output'])['quantity']
        return orchestrator._graph_from_plan({'tasks': {
            'updated': {'tool_id': 'fs_write', 'args': {'path': 'result.txt',
                                                       'content': str(quantity)}},
        }}, goal, tools.TOOL_SPECS, kwargs['completed'])

    monkeypatch.setattr(planner, 'plan', plan)

    async def exercise():
        async def update(graph):
            pass
        with approval.asking(approve):
            return await planner.run(original, update)

    graph = asyncio.run(exercise())
    assert len(calls) == 1
    assert (local / 'result.txt').read_text() == '7'
    assert graph.tasks['read'].output == '{"quantity":7}'
    assert 'write' not in graph.tasks
    assert not graph.failed_tasks()


def test_tool_observations_keep_stable_ids_across_multiple_continuations(local, monkeypatch):
    (local / 'input.txt').write_text('actual input')
    graph = orchestrator._graph_from_plan({'tasks': {
        'TASK_1': {'tool_id': 'fs_read', 'args': {'path': 'input.txt'}},
        'TASK_2': {'tool_id': 'fs_write', 'args': {'path': 'result.txt', 'content': 'draft'},
                   'dependencies': ['TASK_1']},
        'TASK_3': {'tool_id': 'fs_read', 'args': {'path': 'result.txt'},
                   'dependencies': ['TASK_2']},
    }}, 'Copy and verify the input', tools.TOOL_SPECS)
    planner = orchestrator.Orchestrator()
    observed = []

    async def plan(goal, **kwargs):
        if kwargs.get('allow_empty'):
            return orchestrator.ExecutionGraph(goal=goal)
        request = json.loads(goal)
        observed.append([task['id'] for task in request['observations']])
        tasks = {'TASK_3': {'tool_id': 'fs_read', 'args': {'path': 'result.txt'},
                            'dependencies': ['TASK_2']}}
        if len(observed) == 1:
            tasks = {'TASK_2': {'tool_id': 'fs_write', 'args': {
                'path': 'result.txt', 'content': request['observations'][0]['output']},
                'dependencies': ['TASK_1']}, **tasks}
        return orchestrator._graph_from_plan({'tasks': tasks}, goal, tools.TOOL_SPECS,
                                             kwargs['completed'])

    async def update(graph):
        pass

    monkeypatch.setattr(planner, 'plan', plan)

    async def exercise():
        with approval.asking(approve):
            return await planner.run(graph, update, effort='fast')

    result = asyncio.run(exercise())
    assert observed == [['TASK_1'], ['TASK_1', 'TASK_2']]
    assert set(result.tasks) == {'TASK_1', 'TASK_2', 'TASK_3'}
    assert result.tasks['TASK_3'].output == 'actual input'
    assert not result.failed_tasks()


def test_failed_observation_continuation_cannot_execute_stale_writes(local, monkeypatch):
    (local / 'input.txt').write_text('actual data')
    graph = orchestrator._graph_from_plan({'tasks': {
        'read': {'tool_id': 'fs_read', 'args': {'path': 'input.txt'}},
        'write': {'tool_id': 'fs_write', 'args': {'path': 'wrong.txt', 'content': 'guessed'}},
    }}, 'Use the input', tools.TOOL_SPECS)
    planner = orchestrator.Orchestrator()

    async def broken(*args, **kwargs):
        raise RuntimeError('The model did not return a valid continuation')

    monkeypatch.setattr(planner, 'plan', broken)

    async def exercise():
        async def update(graph):
            pass
        with approval.asking(approve):
            await planner.run(graph, update, effort='fast')
    asyncio.run(exercise())
    assert graph.tasks['write'].status == 'failed'
    assert 'tool results' in graph.tasks['write'].error
    assert not (local / 'wrong.txt').exists()


def test_final_review_runs_work_omitted_from_a_shortened_plan(local, monkeypatch):
    graph = orchestrator._graph_from_plan({'tasks': {
        'write': {'tool_id': 'fs_write', 'args': {'path': 'result.txt', 'content': 'beta'}},
    }}, 'Write beta, then calculate 6*7 and verify 42', tools.TOOL_SPECS)
    planner = orchestrator.Orchestrator()
    reviews = []

    async def plan(goal, **kwargs):
        request = json.loads(goal)
        assert kwargs['allow_empty'] and request['task'] == 'verify_completion'
        assert request['complete_output'] == {'tasks': {}}
        reviews.append(request)
        tasks = {} if len(reviews) == 2 else {
            'calculate': {'tool_id': 'python_exec', 'args': {'code': 'print(6*7)'}}}
        if len(reviews) == 2:
            assert json.loads(request['observations'][-1]['output'])['stdout'] == '42\n'
        return orchestrator._graph_from_plan({'tasks': tasks}, goal, tools.TOOL_SPECS,
                                             kwargs['completed'], allow_empty=True)

    async def update(current):
        pass

    monkeypatch.setattr(planner, 'plan', plan)

    async def exercise():
        with approval.asking(approve):
            return await planner.run(graph, update, effort='fast')
    result = asyncio.run(exercise())
    assert len(reviews) == 2
    assert (local / 'result.txt').read_text() == 'beta'
    assert not result.failed_tasks()
    assert list(result.tasks) == ['write', 'calculate']


def test_failed_final_review_never_reports_goal_success(local, monkeypatch):
    graph = orchestrator._graph_from_plan({'tasks': {
        'write': {'tool_id': 'fs_write', 'args': {'path': 'result.txt', 'content': 'beta'}},
    }}, 'Verify the whole goal', tools.TOOL_SPECS)
    planner = orchestrator.Orchestrator()

    async def broken(*args, **kwargs):
        raise RuntimeError('Completion could not be verified')

    async def update(current):
        pass

    monkeypatch.setattr(planner, 'plan', broken)

    async def exercise():
        with approval.asking(approve):
            return await planner.run(graph, update, effort='fast')
    result = asyncio.run(exercise())
    assert (local / 'result.txt').read_text() == 'beta'
    assert result.tasks['write'].status == 'completed'
    assert result.failed_tasks()[0].tool_id == 'completion_check'
    assert 'Completion could not be verified' in result.failed_tasks()[0].error


def test_empty_plan_is_only_valid_for_explicit_completion_review():
    with pytest.raises(ValueError, match='at least one task'):
        orchestrator._graph_from_plan({'tasks': {}}, 'New goal', tools.TOOL_SPECS)
    assert not orchestrator._graph_from_plan({'tasks': {}}, 'Review', tools.TOOL_SPECS,
                                            allow_empty=True).tasks


def test_repeated_observations_are_bounded_without_executing_a_stale_write(local, monkeypatch):
    (local / 'input.txt').write_text('data')
    tasks = {'read': {'tool_id': 'fs_read', 'args': {'path': 'input.txt'}},
             'write': {'tool_id': 'fs_write', 'args': {'path': 'wrong.txt', 'content': 'stale'}}}
    graph = orchestrator._graph_from_plan({'tasks': tasks}, 'Use the data', tools.TOOL_SPECS)
    planner = orchestrator.Orchestrator()
    calls = []

    async def repeat(goal, **kwargs):
        calls.append(goal)
        return orchestrator._graph_from_plan({'tasks': {
            f'again_{len(calls)}': tasks['read'], 'later': tasks['write'],
        }}, goal, tools.TOOL_SPECS, kwargs['completed'])

    monkeypatch.setattr(planner, 'plan', repeat)

    async def exercise():
        async def update(graph):
            pass
        with approval.asking(approve):
            await planner.run(graph, update, effort='fast')
    asyncio.run(exercise())
    assert len(calls) == orchestrator.MAX_ACTIONS
    assert graph.failed_tasks()
    assert any('action limit' in (task.error or '') for task in graph.failed_tasks())
    assert sum(task.status == 'completed' for task in graph.tasks.values()) == (
        orchestrator.MAX_ACTIONS)
    assert not (local / 'wrong.txt').exists()


def test_longer_workflow_finishes_without_the_old_six_observation_ceiling(local, monkeypatch):
    for index in range(8):
        (local / f'{index}.txt').write_text(str(index))
    graph = orchestrator._graph_from_plan({'tasks': {
        'read0': {'tool_id': 'fs_read', 'args': {'path': '0.txt'}},
    }}, 'Combine eight files', tools.TOOL_SPECS)
    planner = orchestrator.Orchestrator()

    async def advance(goal, **kwargs):
        done = kwargs['completed']
        if len(done) < 8:
            tasks = {f'read{len(done)}': {'tool_id': 'fs_read',
                                         'args': {'path': f'{len(done)}.txt'}}}
        elif len(done) == 8:
            tasks = {'combine': {'tool_id': 'fs_write', 'args': {
                'path': 'combined.txt',
                'content': '|'.join(task.output for task in done.values())}}}
        else:
            tasks = {}
        return orchestrator._graph_from_plan({'tasks': tasks}, goal, tools.TOOL_SPECS, done,
                                             allow_empty=kwargs.get('allow_empty', False))
    monkeypatch.setattr(planner, 'plan', advance)

    async def exercise():
        async def update(current):
            pass
        with approval.asking(approve):
            return await planner.run(graph, update, effort='fast')
    result = asyncio.run(exercise())
    assert not result.failed_tasks()
    assert (local / 'combined.txt').read_text() == '0|1|2|3|4|5|6|7'


def test_all_tools_have_typed_contracts_and_reachable_groups():
    exposed = {spec['id'] for spec in tools.tools_for(list(tools.TOOL_GROUPS))}
    for spec in tools.TOOL_SPECS:
        assert spec['id'] in exposed or spec['id'] in {'integrations', 'capability'}
        contract = tool_contract(spec)
        assert contract['parameters']['additionalProperties'] is False
        assert all('type' in field for field in contract['parameters']['properties'].values())
    prompt = orchestrator._tools_description()
    assert all('description' not in item for item in json.loads(prompt))
    assert 'You are' not in prompt


@pytest.mark.skipif(os.name == 'nt', reason='POSIX process-group cancellation')
def test_cancelling_code_stops_the_owned_process(local):
    async def exercise():
        task = asyncio.create_task(execution.execute(cwd=local, code=(
            'import time; from pathlib import Path; Path("started").touch(); '
            'time.sleep(2); Path("finished").touch()')))
        for _ in range(100):
            if (local / 'started').exists():
                break
            await asyncio.sleep(.01)
        assert (local / 'started').exists()
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
        await asyncio.sleep(.05)
        assert not (local / 'finished').exists()
    asyncio.run(exercise())
