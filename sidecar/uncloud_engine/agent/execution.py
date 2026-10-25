"""Local code execution in the engine environment. Permission lives in run_tool."""
from __future__ import annotations

import asyncio
import importlib.util
import json
import os
import signal
import sys
import zipfile
from contextlib import suppress
from pathlib import Path


def environment() -> dict[str, str]:
    env = os.environ.copy()
    env['PATH'] = str(Path(sys.executable).parent) + os.pathsep + env.get('PATH', '')
    env['VIRTUAL_ENV'] = sys.prefix
    env['PYTHONIOENCODING'] = 'utf-8'
    return env


def runtime_info(workspace: Path) -> str:
    return json.dumps({
        'platform': sys.platform, 'python': sys.executable, 'cwd': str(workspace),
        'modules': {name: importlib.util.find_spec(name) is not None for name in (
            'reportlab', 'pypdf', 'docx', 'openpyxl', 'PIL', 'matplotlib')},
    }, separators=(',', ':'))


async def _stop(proc) -> None:
    if proc.returncode is None:
        try:
            if os.name == 'nt':
                killer = await asyncio.create_subprocess_exec(
                    'taskkill', '/PID', str(proc.pid), '/T', '/F',
                    stdout=asyncio.subprocess.DEVNULL, stderr=asyncio.subprocess.DEVNULL)
                await killer.wait()
            else:
                os.killpg(proc.pid, signal.SIGKILL)
        except ProcessLookupError:
            pass
        finally:
            if proc.returncode is None:
                with suppress(ProcessLookupError):
                    proc.kill()
    await proc.communicate()


async def execute(*, cwd: Path, command: str | None = None, code: str | None = None) -> str:
    options = dict(cwd=str(cwd), env=environment(), stdout=asyncio.subprocess.PIPE,
                   stderr=asyncio.subprocess.STDOUT)
    if os.name != 'nt':
        options['start_new_session'] = True
    else:
        import subprocess

        options['creationflags'] = subprocess.CREATE_NEW_PROCESS_GROUP
    if code is not None:
        proc = await asyncio.create_subprocess_exec(sys.executable, '-c', code, **options)
    else:
        proc = await asyncio.create_subprocess_shell(command, **options)
    try:
        stdout, _ = await asyncio.wait_for(proc.communicate(), timeout=120)
    except TimeoutError as exc:
        await _stop(proc)
        raise RuntimeError('Command timed out after 120s') from exc
    except asyncio.CancelledError:
        await _stop(proc)
        raise
    out = stdout.decode('utf-8', errors='replace')[-8000:]
    if proc.returncode != 0:
        raise RuntimeError(f'Exit {proc.returncode}: {out}')
    return out or '(no output)'


def file_info(path: Path) -> dict:
    """Check actual container contents, rather than trusting a filename extension."""
    size = path.stat().st_size
    if not path.is_file() or not size:
        raise ValueError(f'Output is not a nonempty file: {path}')
    suffix = path.suffix.lower()
    if suffix == '.pdf':
        from pypdf import PdfReader

        reader = PdfReader(path)
        if not reader.pages:
            raise ValueError(f'PDF has no pages: {path}')
    elif suffix in {'.docx', '.xlsx', '.pptx'}:
        required = {'.docx': 'word/document.xml', '.xlsx': 'xl/workbook.xml',
                    '.pptx': 'ppt/presentation.xml'}[suffix]
        with zipfile.ZipFile(path) as archive:
            if required not in archive.namelist() or archive.testzip() is not None:
                raise ValueError(f'Invalid {suffix} container: {path}')
    elif suffix in {'.png', '.jpg', '.jpeg', '.webp', '.gif'}:
        from PIL import Image

        with Image.open(path) as image:
            image.verify()
    elif suffix == '.json':
        json.loads(path.read_text(encoding='utf-8'))
    return {'path': str(path), 'bytes': size, 'format': suffix.lstrip('.')}


def read_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == '.pdf':
        from pypdf import PdfReader

        return '\n'.join(page.extract_text() or '' for page in PdfReader(path).pages)[:20000]
    if suffix == '.docx':
        from docx import Document

        document = Document(path)
        text = [paragraph.text for paragraph in document.paragraphs]
        text.extend('\t'.join(cell.text for cell in row.cells)
                    for table in document.tables for row in table.rows)
        return '\n'.join(text)[:20000]
    if suffix == '.xlsx':
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            text = []
            size = 0
            for sheet in workbook:
                text.append(f'[{sheet.title}]')
                for row in sheet.values:
                    line = '\t'.join('' if value is None else str(value) for value in row)
                    text.append(line)
                    size += len(line) + 1
                    if size >= 20000:
                        return '\n'.join(text)[:20000]
            return '\n'.join(text)[:20000]
        finally:
            workbook.close()
    return path.read_text(encoding='utf-8')[:20000]
