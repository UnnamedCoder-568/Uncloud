"""Plain document output from structured content, without generated code."""
from __future__ import annotations

import json
import os
import tempfile
from pathlib import Path
from xml.sax.saxutils import escape

from .execution import file_info, read_file


def validate_document(args: dict) -> None:
    if Path(args['path']).suffix.lower() not in {'.pdf', '.docx', '.xlsx'}:
        raise ValueError('document_write supports .pdf, .docx and .xlsx')
    paragraphs, rows = args.get('paragraphs', []), args.get('rows', [])
    if not paragraphs and not rows:
        raise ValueError('document_write requires paragraphs or rows')
    if len(paragraphs) > 1000 or len(rows) > 1000 or any(len(row) > 30 for row in rows):
        raise ValueError('Document content exceeds the supported page or table size')
    if len(json.dumps(args, ensure_ascii=False)) > 1_000_000:
        raise ValueError('Document content exceeds 1 MB')


def write_document(path: Path, *, title: str = '', paragraphs: list[str] | None = None,
                   rows: list[list] | None = None) -> str:
    paragraphs, rows = paragraphs or [], rows or []
    validate_document({'path': str(path), 'title': title, 'paragraphs': paragraphs, 'rows': rows})
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix='.uncloud-document-', suffix=path.suffix, dir=path.parent)
    os.close(fd)
    temporary = Path(name)
    try:
        suffix = path.suffix.lower()
        if suffix == '.pdf':
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import (
                LongTable,
                Paragraph,
                SimpleDocTemplate,
                Spacer,
                TableStyle,
            )

            styles = getSampleStyleSheet()
            document = SimpleDocTemplate(str(temporary), pagesize=A4, title=title,
                                         leftMargin=48, rightMargin=48,
                                         topMargin=48, bottomMargin=48)
            elements = []
            if title:
                elements += [Paragraph(escape(title), styles['Heading1']), Spacer(1, 12)]
            for paragraph in paragraphs:
                elements += [
                    Paragraph(escape(paragraph).replace('\n', '<br/>'), styles['BodyText']),
                    Spacer(1, 8)]
            if rows:
                columns = max(map(len, rows))
                if not columns:
                    raise ValueError('A document table needs at least one cell')
                cells = [[Paragraph(escape('' if value is None else str(value)), styles['BodyText'])
                          for value in row] + [''] * (columns - len(row)) for row in rows]
                table = LongTable(cells, colWidths=[document.width / columns] * columns,
                                  repeatRows=1)
                table.setStyle(TableStyle([
                    ('GRID', (0, 0), (-1, -1), 0.5, colors.lightgrey),
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#eeeeee')),
                    ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                    ('TOPPADDING', (0, 0), (-1, -1), 6),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                ]))
                elements.append(table)
            document.build(elements)
        elif suffix == '.docx':
            from docx import Document

            document = Document()
            if title:
                document.add_heading(title, 0)
            for paragraph in paragraphs:
                document.add_paragraph(paragraph)
            if rows:
                columns = max(map(len, rows))
                if not columns:
                    raise ValueError('A document table needs at least one cell')
                table = document.add_table(rows=0, cols=columns)
                table.style = 'Table Grid'
                for row in rows:
                    cells = table.add_row().cells
                    for index, value in enumerate(row):
                        cells[index].text = '' if value is None else str(value)
            document.save(str(temporary))
        else:
            from openpyxl import Workbook

            workbook = Workbook()
            workbook.properties.title = title
            sheet = workbook.active
            for paragraph in paragraphs:
                sheet.append([paragraph])
            if paragraphs and rows:
                sheet.append([])
            for row in rows:
                sheet.append(row)
            # Structured strings are data; formulas still require explicit code execution.
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
            workbook.save(temporary)
            workbook.close()
        info = file_info(temporary)
        content = read_file(temporary)
        temporary.replace(path)
        return json.dumps({**info, 'path': str(path), 'text': content}, ensure_ascii=False)
    finally:
        temporary.unlink(missing_ok=True)
